from pathlib import Path
import json
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
INFILE = ROOT / 'build' / 'api_endpoints.json'
OUTFILE = ROOT / 'system_documentation.docx'
USER_MODEL = ROOT / 'src' / 'models' / 'user.js'
README = ROOT / 'README.md'


def safe_read(path: Path):
    if path.exists():
        return path.read_text()
    return ''


def extract_roles(text: str):
    import re
    m = re.search(r"enum:\s*\[(.*?)\]", text, re.S)
    if not m:
        return []
    items = m.group(1)
    return [r.strip().strip("'\"") for r in items.split(',')]


def pick(data, keywords, limit=10):
    out = []
    for e in data:
        for k in keywords:
            if k in e['path']:
                out.append(e)
                break
        if len(out) >= limit:
            break
    return out


def main():
    doc = Document()
    doc.add_heading('System Documentation - User Workflows', level=1)

    readme = safe_read(README)
    if readme:
        doc.add_heading('Project Overview', level=2)
        doc.add_paragraph(readme.strip().split('\n\n',1)[0])

    user_text = safe_read(USER_MODEL)
    roles = extract_roles(user_text)
    if 'client' not in roles:
        roles.insert(0,'client')

    data = []
    if INFILE.exists():
        data = json.loads(INFILE.read_text())

    # Define workflows
    workflows = {
        'client': {
            'title': 'Client (End user)',
            'steps': [
                'Search rooms and filter by date, price, facilities.',
                'View room/accommodation details and availability.',
                'Create booking request with guest details.',
                'Pay via chosen provider (Chapa/Stripe) or select pay-on-site.',
                'Receive confirmation email and notification.',
                'View booking history and rate stay after checkout.'
            ],
            'keywords': ['filter','booking','request','history','pay','chapa']
        },
        'owner': {
            'title': 'Owner',
            'steps': [
                'Log into owner dashboard (web).',
                'Create/edit accommodations, rooms, room types and facilities.',
                'Set prices, manage availability and promotions.',
                'Accept or decline booking requests and manage check-ins.',
                'Connect bank accounts and view profit reports.'
            ],
            'keywords': ['accommodation','rooms','price','profit','bank']
        },
        'receptionist': {
            'title': 'Receptionist',
            'steps': [
                'View daily arrivals and departures.',
                'Confirm guest check-ins and check-outs.',
                'Handle in-person payments and issue receipts.',
                'Modify bookings per guest requests.'
            ],
            'keywords': ['booking','checkin','turnToCheckin','payment']
        },
        'sales': {
            'title': 'Sales',
            'steps': [
                'Assist potential customers with availability.',
                'Create bookings and apply packages or promotions.',
                'Access sales reports and export data.'
            ],
            'keywords': ['package','booking','profit']
        },
        'call_center': {
            'title': 'Call Center Agent',
            'steps': [
                'Search clients by phone/email.',
                'Create bookings on behalf of clients.',
                'Process cancellations/refunds or escalate to owners.'
            ],
            'keywords': ['client','booking','refund','cancel']
        },
        'super_admin': {
            'title': 'Super Admin',
            'steps': [
                'Manage users and roles.',
                'Change global settings and integration keys.',
                'Access system-wide reports and logs.'
            ],
            'keywords': ['user','settings','webhook','payment']
        }
    }

    doc.add_heading('Detailed Workflows by Role', level=2)
    for r in roles:
        key = r if r in workflows else None
        if not key:
            title = r.replace('_',' ').title()
            steps = ['Authenticate as role.','Perform role-specific tasks via UI or API.']
            keywords = [r]
        else:
            title = workflows[key]['title']
            steps = workflows[key]['steps']
            keywords = workflows[key]['keywords']

        doc.add_heading(title, level=3)
        doc.add_paragraph('Summary of responsibilities:')
        for s in steps:
            doc.add_paragraph('- ' + s)

        if data:
            doc.add_paragraph('Representative API endpoints:')
            picks = pick(data, keywords, limit=10)
            if picks:
                for e in picks:
                    doc.add_paragraph(f"- {e['method']} {e['path']}  —  {e['file']}:{e['line']}")
            else:
                doc.add_paragraph('- No matching endpoints found in extraction; update keywords if necessary.')

        doc.add_paragraph('Authentication & notes:')
        doc.add_paragraph('- Endpoints protected by Firebase auth or role-based `accessControl`.')
        doc.add_paragraph('- For user manuals, include UI screenshots and sample request/response payloads.')

    doc.save(OUTFILE)
    print('Wrote', OUTFILE)

if __name__=='__main__':
    main()
