import json
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
INFILE = ROOT / 'build' / 'api_endpoints.json'
OUTFILE = ROOT / 'system_documentation.docx'
README = ROOT / 'README.md'
USER_MODEL = ROOT / 'src' / 'models' / 'user.js'


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def safe_read(path: Path) -> str:
    if path.exists():
        return path.read_text()
    return ''


def gather_summary(data):
    summary = {}
    for e in data:
        p = e['path']
        if 'booking' in p:
            summary.setdefault('booking', []).append(e)
        if 'payment' in p or 'chapa' in p or 'stripe' in p:
            summary.setdefault('payment', []).append(e)
        if 'client' in p or 'auth' in p or 'user' in p:
            summary.setdefault('client', []).append(e)
        if 'notification' in p or 'notify' in p:
            summary.setdefault('notification', []).append(e)
    return summary


def extract_roles(user_model_text: str):
    # naive extraction of enum values
    import re
    m = re.search(r"enum:\s*\[(.*?)\]", user_model_text, re.S)
    if not m:
        return []
    items = m.group(1)
    # split by comma and strip quotes/spaces
    roles = [r.strip().strip("'\"") for r in items.split(',')]
    return roles


def pick_endpoints_for_keywords(data, keywords, limit=8):
    picks = []
    for e in data:
        for kw in keywords:
            if kw in e['path']:
                picks.append(e)
                break
        if len(picks) >= limit:
            break
    return picks


def add_role_workflows(doc, data, roles):
    doc.add_heading('User Role Workflows', level=2)
    # Define workflows for each role
    role_defs = {
        'client': {
            'display': 'Client (End user)',
            'steps': [
                'Search available rooms and filters (availability, price, facilities).',
                'View room and accommodation details.',
                'Create booking request; provide guest details and payment method.',
                'Complete payment via Chapa/Stripe (redirect or direct charge).',
                'Receive booking confirmation and view booking history.',
                'Request cancellations or rate the stay after checkout.'
            ],
            'keywords': ['filter', 'request', 'booking', 'history', 'pay', 'chapa']
        },
        'owner': {
            'display': 'Owner (Accommodation manager)',
            'steps': [
                'Log in to owner dashboard.',
                'Create and update accommodations, rooms, and room types.',
                'Set prices and availability; change nightly/daily prices.',
                'Accept or reject booking requests and view pending bookings.',
                'Manage bank accounts for payouts and view profits.'
            ],
            'keywords': ['accommodation', 'rooms', 'booking', 'price', 'profit', 'bank']
        },
        'receptionist': {
            'display': 'Receptionist',
            'steps': [
                'View today's check-ins and check-outs.',
                'Confirm guest arrivals and change booking statuses (check-in/check-out).',
                'Process on-site payments and issue receipts.',
                'Assist clients with booking modifications.'
            ],
            'keywords': ['booking', 'checkin', 'turnToCheckin', 'payment']
        },
        'sales': {
            'display': 'Sales',
            'steps': [
                'Search bookings and clients to assist sales-related inquiries.',
                'Manage promotions or packages and coordinate with owners.',
                'Generate reports for revenue and conversions.'
            ],
            'keywords': ['package', 'booking', 'report', 'profit']
        },
        'call_center': {
            'display': 'Call Center Agent',
            'steps': [
                'Lookup client profiles by phone/email.',
                'Create bookings on the client\'s behalf via request endpoints.',
                'Handle basic refunds, cancellations, and escalate to owners when required.'
            ],
            'keywords': ['client', 'booking', 'refund', 'cancel']
        },
        'super_admin': {
            'display': 'Super Admin',
            'steps': [
                'Manage all users and roles (create admins, change roles).',
                'Access system-wide reports and settings.',
                'Manage global configurations and integrations (payment keys, webhooks).'
            ],
            'keywords': ['user', 'settings', 'payment', 'webhook']
        }
    }

    # Ensure clients role is present in roles list
    all_roles = roles[:] if roles else []
    if 'client' not in all_roles:
        all_roles.insert(0, 'client')

    for r in all_roles:
        key = r
        if r not in role_defs:
            # generic role template
            role_title = r.replace('_', ' ').title()
            role_def = {
                'display': role_title,
                'steps': ['Login / authenticate as role.', 'Perform role-specific actions via API or dashboard.'],
                'keywords': [r]
            }
        else:
            role_def = role_defs[r]

        doc.add_heading(role_def['display'], level=3)
        doc.add_paragraph('Primary tasks:')
        for s in role_def['steps']:
            doc.add_paragraph('- ' + s)

        # Show example endpoints
        picks = pick_endpoints_for_keywords(data, role_def['keywords'], limit=10)
        if picks:
            doc.add_paragraph('Representative API endpoints:')
            for e in picks:
                doc.add_paragraph(f"- {e['method']} {e['path']}  —  {e['file']}:{e['line']}")
        doc.add_paragraph('Authentication:')
        doc.add_paragraph('- Most protected endpoints require Firebase auth tokens or role-based access via `accessControl`.')
        doc.add_paragraph('Notes:')
        doc.add_paragraph('- Replace example endpoints with concrete frontend routes and payload examples when preparing user manuals.')


def main():
    doc = Document()
    doc.add_heading('System Documentation', level=1)

    # Project overview from README
    readme = safe_read(README)
    if readme:
        doc.add_heading('Project Overview', level=2)
        # take the first paragraph
        first_para = readme.strip().split('\n\n', 1)[0]
        doc.add_paragraph(first_para)

    # Installation & running
    doc.add_heading('Installation & Run', level=2)
    doc.add_paragraph('Requirements: Node >= 22.21.0, MongoDB. See README for full details.')
    doc.add_paragraph('Quickstart commands:')
    doc.add_paragraph('npm install')
    doc.add_paragraph('npm start')

    # Environment and scripts
    if readme:
        doc.add_heading('Environment Variables', level=3)
        # find the Environment variables section if present
        if 'Environment variables' in readme:
            part = readme.split('Environment variables', 1)[1]
            env_lines = part.split('\n')[:20]
            for l in env_lines:
                if l.strip().startswith('-'):
                    doc.add_paragraph(l.strip().lstrip('- ').strip())
                else:
                    break

    # API endpoints
    doc.add_heading('API Endpoints (summary)', level=2)
    if not INFILE.exists():
        doc.add_paragraph('No API endpoints file found. Run the extractor script first.')
        doc.save(OUTFILE)
        print('Wrote', OUTFILE)
        return

    data = json.loads(INFILE.read_text())
    summary = gather_summary(data)

    # Roles
    user_model = safe_read(USER_MODEL)
    roles = extract_roles(user_model)
    if roles:
        doc.add_heading('User Roles', level=3)
        doc.add_paragraph('The application defines these user roles: ' + ', '.join(roles))

    # Workflows
    doc.add_heading('Primary User Workflows', level=2)
    # Booking flow
    if 'booking' in summary:
        doc.add_heading('1) Booking Flow (Client)', level=3)
        doc.add_paragraph('Typical sequence:')
        doc.add_paragraph('- Client searches for rooms using POST /filter or /old-filter')
        doc.add_paragraph('- Client requests a booking via POST /request')
        doc.add_paragraph('- System creates a booking record and may require payment')
        doc.add_paragraph('- Client can view booking history via GET /history')
        doc.add_paragraph('Key endpoints:')
        for e in summary['booking'][:10]:
            doc.add_paragraph(f"- {e['method']} {e['path']} (defined in {e['file']}:{e['line']})")

    # Payment flow
    if 'payment' in summary:
        doc.add_heading('2) Payment Flow', level=3)
        doc.add_paragraph('Typical sequence:')
        doc.add_paragraph('- After booking, client is redirected to payment endpoints (Chapa/Stripe)')
        doc.add_paragraph('- Webhook endpoints handle asynchronous payment confirmations')
        doc.add_paragraph('Key endpoints:')
        for e in summary['payment'][:12]:
            doc.add_paragraph(f"- {e['method']} {e['path']} (defined in {e['file']}:{e['line']})")

    # Notification flow
    if 'notification' in summary:
        doc.add_heading('3) Notification Flow', level=3)
        doc.add_paragraph('- Notifications are sent via Firebase and internal jobs')
        doc.add_paragraph('Key endpoints:')
        for e in summary['notification'][:8]:
            doc.add_paragraph(f"- {e['method']} {e['path']} (defined in {e['file']}:{e['line']})")

    # Client & Auth
    if 'client' in summary:
        doc.add_heading('Client & Authentication', level=2)
        doc.add_paragraph('Clients authenticate via Firebase tokens or local auth depending on endpoint.')
        doc.add_paragraph('Key client endpoints:')
        for e in summary['client'][:12]:
            doc.add_paragraph(f"- {e['method']} {e['path']} (defined in {e['file']}:{e['line']})")

    # Full endpoint listing grouped by file (compact)
    doc.add_page_break()
    doc.add_heading('Full Endpoint Listing', level=1)
    files = {}
    for e in data:
        files.setdefault(e['file'], []).append(e)

    for file, endpoints in sorted(files.items()):
        add_heading(doc, file, level=2)
        for ep in endpoints:
            doc.add_paragraph(f"{ep['method']} {ep['path']}  —  defined in {ep['file']}:{ep['line']}")

    doc.add_page_break()
    doc.add_heading('Notes', level=2)
    doc.add_paragraph('This document is an auto-generated starting point. Add diagrams, sequence examples, request/response schemas, and security details for production-ready documentation.')

    doc.save(OUTFILE)
    print('Wrote', OUTFILE)


if __name__ == '__main__':
    main()
